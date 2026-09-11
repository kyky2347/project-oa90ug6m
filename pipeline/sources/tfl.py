import json
import re
import xml.etree.ElementTree as ET
from collections import defaultdict
from difflib import get_close_matches
from urllib.parse import quote

import numpy as np
import openpyxl
from shapely import Point

from pipeline.sources.base import Asset, Source, request

BUCKET = "https://s3-eu-west-1.amazonaws.com/crowding.data.tfl.gov.uk/"
NS = {"s": "http://s3.amazonaws.com/doc/2006-03-01/"}


def normal_name(name):
    s = re.sub(
        r"\b(underground|overground|rail|railway|station|dlr|lu|nr|elizabeth line|tram stop)\b",
        "",
        name.lower(),
    )
    s = s.replace("&", "and").replace("st.", "st").replace("saint ", "st ")
    return re.sub("[^a-z0-9]", "", s)


def parse_flows(workbook):
    output = {}
    for sheet in ["Station_Entries", "Station_Exits"]:
        if sheet not in workbook.sheetnames:
            raise ValueError(f"TfL sheet missing: {sheet}")
        it = workbook[sheet].iter_rows(values_only=True)
        header = None
        for r in it:
            if r[0] == "NLC":
                header = r
                break
        if not header or "Station" not in header or "ASC" not in header:
            raise ValueError("TfL station schema changed")
        periods = [
            (i, int(v[:2]) % 24 * 4 + int(v[2:4]) // 15)
            for i, v in enumerate(header)
            if isinstance(v, str) and re.match(r"^\d{4}-\d{4}$", v)
        ]
        if len(periods) != 96:
            raise ValueError("Expected 96 TfL quarter-hour periods")
        for row in it:
            if not row[0] or not row[1] or not isinstance(row[2], str):
                continue
            sid = str(row[1])
            record = output.setdefault(sid, {"id": sid, "name": row[2], "values": np.zeros(96)})
            for idx, band in periods:
                v = row[idx]
                if not isinstance(v, (int, float)) or not np.isfinite(v) or v < 0:
                    raise ValueError("Missing/invalid TfL timeband")
                record["values"][band] += v
    if not output:
        raise ValueError("No station flows")
    return output


class TfLSource(Source):
    id = "tfl"
    name = "NUMBAT typical station entries and exits"
    publisher = "Transport for London"
    url = "https://tfl.gov.uk/info-for/open-data-users/our-open-data"
    license = "TfL transport data terms and conditions (open data)"

    def discover(self):
        entries, token = [], None
        while True:
            params = {"list-type": 2, "prefix": "NUMBAT/"}
            if token:
                params["continuation-token"] = token
            tree = ET.fromstring(request(BUCKET, params=params).content)
            entries.extend(
                (e.find("s:Key", NS).text, e.find("s:LastModified", NS).text)
                for e in tree.findall("s:Contents", NS)
            )
            nxt = tree.find("s:NextContinuationToken", NS)
            if nxt is None:
                break
            token = nxt.text
        files = [(k, date) for k, date in entries if re.search(r"NBT\d{2}(TWT|SAT|SUN)_Outputs\.xlsx$", k)]
        if not files:
            raise ValueError("No official NUMBAT time-profile workbooks found")
        year = max(int(re.search(r"NUMBAT (\d{4})/", k).group(1)) for k, _ in files)
        out = []
        for k, modified in files:
            if f"NUMBAT {year}/" in k:
                day = {"TWT": "weekday", "SAT": "saturday", "SUN": "sunday"}[
                    re.search(r"NBT\d{2}(TWT|SAT|SUN)", k).group(1)
                ]
                out.append(
                    Asset(
                        "https://crowding.data.tfl.gov.uk/" + quote(k),
                        day + ".xlsx",
                        f"{year}-11-01",
                        {"year": year, "day": day, "released_at": modified},
                    )
                )
        out.append(
            Asset("https://api.tfl.gov.uk/StopPoint/Mode/tube,dlr,overground,elizabeth-line", "stops.json")
        )
        return sorted(out, key=lambda a: a.filename)

    def transform(self, directory, assets):
        stops = json.loads((directory / "stops.json").read_text())
        if "stopPoints" not in stops:
            raise ValueError("TfL stop geometry schema changed")
        lookup = defaultdict(list)
        for s in stops["stopPoints"]:
            if -0.65 < s.get("lon", 9) < 0.4 and 51.25 < s.get("lat", 0) < 51.75:
                lookup[normal_name(s["commonName"])].append((s["lon"], s["lat"]))
        stations, profiles, unmatched = {}, [], set()
        for a in assets:
            if not a.filename.endswith(".xlsx"):
                continue
            w = openpyxl.load_workbook(directory / a.filename, read_only=True, data_only=True)
            flows = parse_flows(w)
            w.close()
            for sid, row in flows.items():
                key = normal_name(row["name"])
                quality = 1.0
                if key not in lookup:
                    found = get_close_matches(key, lookup, n=1, cutoff=0.94)
                    if not found:
                        unmatched.add(row["name"])
                        continue
                    key, quality = found[0], 0.8
                lon, lat = np.median(lookup[key], axis=0)
                total = float(np.sum(row["values"]))
                if total <= 0:
                    continue
                if a.extra["day"] == "weekday" or sid not in stations:
                    stations[sid] = dict(
                        id=sid,
                        name=row["name"],
                        geom=Point(lon, lat).wkt,
                        daily_demand=total,
                        match_quality=quality,
                    )
                profiles.append(dict(station_id=sid, day_type=a.extra["day"], values=row["values"].tolist()))
        if len(stations) < 300:
            raise ValueError(f"Insufficient geocoded TfL stations: {len(stations)}")
        profiles = [p for p in profiles if p["station_id"] in stations]
        self.notes = [
            "Typical autumn station entries + exits; weekday is Tuesday–Thursday. Not live footfall.",
            "Station locations matched to official TfL StopPoint names; fuzzy matches lower spatial confidence.",
            f"Unmatched/out-of-study-area station names ({len(unmatched)}): " + ", ".join(sorted(unmatched)),
        ]
        return {"transport_stations": list(stations.values()), "transport_profiles": profiles}
