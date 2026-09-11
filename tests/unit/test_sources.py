import hashlib
import json
from datetime import UTC, datetime
from pathlib import Path
from unittest.mock import Mock

import geopandas as gpd
import httpx
import numpy as np
import openpyxl
import pytest
from shapely import Point, Polygon

from pipeline.sources import base
from pipeline.sources.base import Asset, Source, download_asset
from pipeline.sources.geography import valid_multi, validate_geo
from pipeline.sources.ons import parse_population
from pipeline.sources.osm import POIHandler, taxonomy
from pipeline.sources.police import parse_incidents
from pipeline.sources.registry import SOURCES
from pipeline.sources.tfl import normal_name, parse_flows
from pipeline.sources.voa import grouped_median, parse_band_csv

FIXTURES = Path(__file__).parents[2] / "data/fixtures"


def test_ons_real_fixture_and_schema_contract():
    rows = json.loads((FIXTURES / "ons.json").read_text())
    result = parse_population(rows)
    assert len(result) == len(rows) and all(r["population"] > 0 for r in result)
    with pytest.raises(ValueError):
        parse_population([{"LSOA 2021 Code": "E01000001"}])
    with pytest.raises(ValueError):
        parse_population([])
    with pytest.raises(ValueError):
        parse_population([{**rows[0], "Total": -3}])


def test_geography_and_gla_real_polygon_contracts():
    fixture = json.loads((FIXTURES / "polygons.geojson").read_text())
    frame = gpd.GeoDataFrame.from_features(fixture, crs=4326)
    assert len(validate_geo(frame)) == 2
    assert all(valid_multi(g).is_valid for g in frame.geometry)
    with pytest.raises(ValueError):
        valid_multi(Point(0, 0))
    with pytest.raises(ValueError):
        validate_geo(gpd.GeoDataFrame(geometry=[Polygon([(1, 1), (2, 1), (2, 2), (1, 1)])], crs=4326))
    with pytest.raises(ValueError):
        validate_geo(gpd.GeoDataFrame(geometry=[], crs=4326))


@pytest.mark.parametrize(
    "tags,category",
    [
        ({"amenity": "cafe", "shop": "bakery"}, "coffee"),
        ({"office": "coworking"}, "coworking"),
        ({"leisure": "fitness_centre"}, "gym_fitness"),
        ({"shop": "vacant"}, None),
        ({"amenity": "school"}, "school"),
        ({"railway": "station"}, "transport"),
    ],
)
def test_osm_taxonomy_single_assignment(tags, category):
    assert taxonomy(tags) == category


def test_osm_handler_real_fixture_and_london_filter():
    handler = POIHandler()
    item = json.loads((FIXTURES / "osm.json").read_text())
    handler.add(item["id"], item["tags"], Point(item["longitude"], item["latitude"]))
    handler.add("outside", {"amenity": "cafe"}, Point(2, 48))
    assert len(handler.records) == 1
    assert next(iter(handler.records.values()))["h3_8"].startswith("88")


def test_tfl_real_quarter_hours_and_invalid_period():
    wb = openpyxl.load_workbook(FIXTURES / "tfl.xlsx", data_only=True)
    data = parse_flows(wb)
    assert len(data) == 1
    assert len(next(iter(data.values()))["values"]) == 96
    assert np.sum(next(iter(data.values()))["values"]) > 0
    wb["Station_Entries"].cell(2, 14).value = -1
    with pytest.raises(ValueError):
        parse_flows(wb)
    assert normal_name("Baker Street Underground Station") == normal_name("Baker Street LU")
    with pytest.raises(ValueError):
        parse_flows(openpyxl.Workbook())


def test_voa_real_band_and_suppression():
    result = parse_band_csv((FIXTURES / "voa.csv").read_bytes())
    assert len(result) == 2
    assert grouped_median([100, 0, 0, 0, 0]) == 6000
    assert grouped_median([0, 0, 100, 0, 0]) == 33000
    assert grouped_median([None, 10, 20, 30, 40]) is None
    assert grouped_median([0, 0, 0, 0, 100]) is None
    assert grouped_median([0] * 5) is None
    with pytest.raises(ValueError):
        parse_band_csv(b"unknown\nfield\n")


def test_police_approximate_real_fixture():
    records = json.loads((FIXTURES / "police.json").read_text())
    assert len(parse_incidents(records)) == 2
    assert parse_incidents([{**records[0], "location": None}]) == []
    assert parse_incidents([]) == []
    with pytest.raises(ValueError):
        parse_incidents({"error": "changed"})
    with pytest.raises(ValueError):
        parse_incidents([{"id": 1}])


def mock_http(monkeypatch, handler):
    monkeypatch.setattr(base, "client", lambda: httpx.Client(transport=httpx.MockTransport(handler)))
    monkeypatch.setattr(base.time, "sleep", lambda _: None)


def test_streamed_download_checksum_atomic_retry_and_cache(monkeypatch, tmp_path):
    calls = []

    def response(request):
        calls.append(request.method)
        return httpx.Response(
            503 if len(calls) == 1 else 200,
            content=b"official fixture",
            headers={"content-type": "application/octet-stream", "etag": "v1"},
        )

    mock_http(monkeypatch, response)
    target = tmp_path / "file.bin"
    asset = Asset(
        "https://example.org/file.bin",
        "file.bin",
        extra={"sha256": hashlib.sha256(b"official fixture").hexdigest()},
    )
    receipt = download_asset(asset, target)
    assert target.read_bytes() == b"official fixture" and len(calls) == 2
    assert receipt["etag"] == "v1" and not target.with_suffix(".bin.part").exists()
    monkeypatch.setattr(base, "ROOT", tmp_path)
    Source().download([asset], tmp_path / "raw1")
    n = len(calls)
    Source().download([asset], tmp_path / "raw2")
    assert len(calls) == n and (tmp_path / "raw2/file.bin").read_bytes() == target.read_bytes()


@pytest.mark.parametrize(
    "body,headers,max_bytes,checksum",
    [
        (b"<html>error", {"content-type": "text/html"}, 100, None),
        (b"", {}, 100, None),
        (b"12345", {}, 3, None),
        (b"abc", {}, 100, "wrong"),
    ],
)
def test_reject_invalid_downloads(monkeypatch, tmp_path, body, headers, max_bytes, checksum):
    mock_http(monkeypatch, lambda r: httpx.Response(200, content=body, headers=headers))
    a = Asset("https://example.org/x", "x.zip", extra={"sha256": checksum})
    with pytest.raises(ValueError):
        download_asset(a, tmp_path / "x.zip", max_bytes)
    assert not (tmp_path / "x.zip").exists() and not (tmp_path / "x.zip.part").exists()


def test_source_updates_use_upstream_validators(monkeypatch):
    a = Asset("https://example.org/x", "x.csv", "2026-01-01")
    old = {
        "id": "old",
        "parser_version": "1.0.0",
        "retrieved_at": datetime.now(UTC),
        "metadata": {
            "assets": [{"url": a.url, "published_at": a.published_at, "download": {"etag": "same"}}]
        },
    }
    mock_http(monkeypatch, lambda r: httpx.Response(200, headers={"etag": "same"}))
    assert Source().check_for_update([a], old) is False
    mock_http(monkeypatch, lambda r: httpx.Response(200, headers={"etag": "new"}))
    assert Source().check_for_update([a], old) is True


def test_failure_preserves_verified_snapshot_and_records_failure(monkeypatch):
    source = Source()
    source.id = "test"
    source.register = Mock()
    source.previous = Mock(return_value={"id": "verified-before"})
    source.discover = Mock(side_effect=ValueError("schema changed"))
    execute = Mock()
    monkeypatch.setattr(base, "execute", execute)
    assert source.ingest() == "verified-before"
    assert execute.call_args[0][1]["error"] == "schema changed"


def test_every_adapter_has_complete_registry_contract(tmp_path):
    assert len(SOURCES) == 7
    for source in SOURCES.values():
        assert all(
            callable(getattr(source, k))
            for k in (
                "discover",
                "check_for_update",
                "download",
                "validate_raw",
                "transform",
                "load",
                "metadata",
            )
        )
        assert all(source.metadata()[k] for k in ("id", "name", "publisher", "url", "license"))
        with pytest.raises(ValueError):
            source.validate_raw(tmp_path, [Asset("https://example.org/missing", "missing.zip")])


@pytest.mark.parametrize(
    "url, expected",
    [
        ("https://www.ons.gov.uk/mid2022revisednov2025tomid2024/sapelsoasyoa20222024.xlsx", 2024),
        ("https://www.ons.gov.uk/published2027/population20252026.xlsx", 2026),
    ],
)
def test_population_reference_ignores_publication_year(url, expected):
    from pipeline.sources.ons import population_reference_year

    assert population_reference_year(url) == expected
